"""Generate Excel report for Selenium WebDriver test suite.
Output: reports/selenium_report.xlsx
Run: python scripts/generate_report.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "reports",
    "selenium_report.xlsx",
)

# === Test data (from actual pytest run) ===
# Tuple: (TT_ID, function, objective, input, expected, tc_type,
#         check_ui, check_db, rollback, notes, result)
TESTS = [
    ("TT-S001", "test_TC001_valid_login",
     "Student login with valid credentials -> redirect away from /login",
     "student_code='SV_TEST_01', password='NewTest@12345'",
     "URL no longer contains /login (navigated to /student)",
     "Positive", "URL redirect", "N", "N",
     "Page Object Model (LoginPage); Selenium Manager auto-handles Edge driver",
     "PASS"),

    ("TT-S002", "test_TC002_wrong_password",
     "Wrong password -> error message displayed",
     "student_code='SV_TEST_01', password='wrong_pass_xxx'",
     "Error message visible AND URL still on /login",
     "Negative", "Error text", "N", "N",
     "WebDriverWait + visibility of '.bg-red-50.text-red-600'",
     "PASS"),

    ("TT-S003", "test_TC003_nonexistent_code",
     "Non-existent student code -> error message",
     "student_code='B99999_NOT_EXIST', password='any_pass'",
     "Error message visible AND URL still on /login",
     "Negative", "Error text", "N", "N",
     "Backend returns generic 401 -> security best practice (no user enumeration)",
     "PASS"),

    ("TT-S004", "test_TC004_missing_student_code",
     "Empty student code field -> validation blocks submit",
     "Only password='any_pass', student_code is empty",
     "Form not submitted, URL still /student/login",
     "Negative", "Validation msg", "N", "N",
     "HTML5 'required' attribute prevents submission",
     "PASS"),

    ("TT-S005", "test_TC005_missing_password",
     "Empty password field -> validation blocks submit",
     "Only student_code='SV_TEST_01', password is empty",
     "Form not submitted, URL still /student/login",
     "Negative", "Validation msg", "N", "N",
     "HTML5 'required' attribute prevents submission",
     "PASS"),

    ("TT-S006", "test_TC006_valid_admin_login",
     "Admin login with valid credentials -> redirect away from /admin/login",
     "email='admin_test@cntt1.edu.vn', password='Admin@2024'",
     "URL no longer contains /login (navigated to /admin)",
     "Positive", "URL redirect", "N", "N",
     "Reuses LoginPage with enter_email() instead of enter_student_code()",
     "PASS"),

    ("TT-S007", "test_TC007_nonexistent_email",
     "Admin email does not exist -> error message",
     "email='notexist_xxx@ptit.edu.vn', password='any_pass'",
     "Error message visible AND URL still on /admin/login",
     "Negative", "Error text", "N", "N",
     "Page Object LoginPage (URL_ADMIN)",
     "PASS"),

    ("TT-S008", "test_TC008_wrong_admin_password",
     "Admin wrong password -> error message",
     "email='admin_test@cntt1.edu.vn', password='wrong_password_xxx'",
     "Error message visible AND URL still on /admin/login",
     "Negative", "Error text", "N", "N",
     "Backend returns generic 'Invalid credentials'",
     "PASS"),

    ("TT-S009", "test_TC013_change_password_success",
     "Change password success + CheckDB + Rollback",
     "current='NewTest@12345', new='NewPass@123', confirm='NewPass@123'",
     "UI shows success message; DB password_hash differs from original",
     "Positive", "Success msg", "password_hash", "rollback_password() in finally",
     "Fixture driver_logged_in; READ COMMITTED + autocommit; try/finally guarantees rollback",
     "PASS"),

    ("TT-S010", "test_TC014_wrong_current_password",
     "Wrong current password -> error message at change-password page",
     "current='WRONG_OLD_PASS', new='NewPass@123', confirm='NewPass@123'",
     "Error message visible at /student/change-password",
     "Negative", "Error text", "N", "N",
     "[BUG #1] Backend returns 401 -> axios interceptor (api.js:33-50) auto-logs out user. "
     "See BUGS.md for fix proposals.",
     "FAIL"),

    ("TT-S011", "test_TC016_password_too_short",
     "New password < 6 characters -> error message",
     "current='NewTest@12345', new='abc', confirm='abc'",
     "Error visible (React error OR HTML5 validationMessage)",
     "Negative", "Error text", "N", "N",
     "Boundary test: length=3 < 6. Checks both error sources via execute_script()",
     "PASS"),

    ("TT-S012", "test_TC017_confirm_mismatch",
     "Confirm password does not match -> error message",
     "current='NewTest@12345', new='NewPass@123', confirm='DifferentPass99'",
     "Error message visible ('Confirm password does not match')",
     "Negative", "Error text", "N", "N",
     "Frontend cross-field validation in handleSubmit()",
     "PASS"),

    ("TT-S013", "test_login_data_driven[td0]",
     "[Data-driven row 0] Valid student login from Excel sheet",
     "Excel row: code='SV_TEST_01', pass='NewTest@12345', expected='success'",
     "Login successful -> redirected away from /login",
     "Positive", "URL redirect", "N", "N",
     "@pytest.mark.parametrize + openpyxl.load_workbook(); reads StudentLogin sheet",
     "PASS"),

    ("TT-S014", "test_login_data_driven[td1]",
     "[Data-driven row 1] Wrong password from Excel",
     "Excel row: code='SV_TEST_01', pass='wrong_pass', expected='khong dung'",
     "Error message visible AND URL still on /login",
     "Negative", "Error text", "N", "N",
     "Parametrized test reuses LoginPage; expected='khong dung' -> branch checking error",
     "PASS"),

    ("TT-S015", "test_login_data_driven[td2]",
     "[Data-driven row 2] Non-existent code from Excel",
     "Excel row: code='B99999_NOT_EXIST', pass='any_pass', expected='khong dung'",
     "Error message visible AND URL still on /login",
     "Negative", "Error text", "N", "N",
     "Parametrized; same flow as TT-S014 with different input",
     "PASS"),

    ("TT-S016", "test_login_data_driven[td3]",
     "[Data-driven row 3] Empty student code from Excel",
     "Excel row: code='', pass='pass', expected='validation'",
     "URL still on /student/login (HTML5 required blocks submit)",
     "Negative", "Validation msg", "N", "N",
     "Excel cell None converted to '' in load_login_data(); branch checking URL unchanged",
     "PASS"),

    ("TT-S017", "test_login_data_driven[td4]",
     "[Data-driven row 4] Empty password from Excel",
     "Excel row: code='SV_TEST_01', pass='', expected='validation'",
     "URL still on /student/login (HTML5 required blocks submit)",
     "Negative", "Validation msg", "N", "N",
     "Parametrized; demonstrates data-driven coverage of TC005",
     "PASS"),
]

# === Styling ===
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="2E75B6")
SUBTITLE_FILL = PatternFill("solid", fgColor="9DC3E6")
PASS_FILL = PatternFill("solid", fgColor="C6E0B4")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
POSITIVE_FILL = PatternFill("solid", fgColor="C6E0B4")
NEGATIVE_FILL = PatternFill("solid", fgColor="FFD9B3")
MIXED_FILL = PatternFill("solid", fgColor="FFF2CC")

WHITE_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True, size=10)
NORMAL = Font(size=10)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN = Side(border_style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_cell(cell, fill=None, font=None, alignment=None, border=BORDER):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    cell.border = border


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Selenium Report"

    headers = [
        "TT ID", "Tool", "Function Name / Script", "Testing Objective",
        "Input", "Expected Output", "TC Type", "Check UI", "CheckDB",
        "Rollback", "Notes / Techniques", "Result"
    ]
    n_cols = len(headers)

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title = ws.cell(row=1, column=1,
                    value="PART 1: SELENIUM WEBDRIVER - Dinh Nghia (Auth & Password)")
    style_cell(title, fill=TITLE_FILL,
               font=Font(bold=True, color="FFFFFF", size=14),
               alignment=CENTER)
    ws.row_dimensions[1].height = 28

    # Row 2: Subtitle
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub = ws.cell(row=2, column=1,
                  value="Based on: TC001-TC020 (Unit Test) | "
                        "Tool: Python Selenium WebDriver 4.x + pytest + openpyxl + "
                        "mysql-connector + bcrypt | "
                        "Browser: Microsoft Edge (Selenium Manager)")
    style_cell(sub, fill=SUBTITLE_FILL, font=Font(bold=True, size=10),
               alignment=CENTER)
    ws.row_dimensions[2].height = 22

    # Row 3: Headers
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        style_cell(cell, fill=HEADER_FILL, font=WHITE_FONT, alignment=CENTER)
    ws.row_dimensions[3].height = 30

    # Data rows
    pass_count = 0
    fail_count = 0
    positive_count = 0
    negative_count = 0
    mixed_count = 0

    for i, t in enumerate(TESTS, start=4):
        (tt_id, fn, obj, inp, exp, tc_type, ui, db, rb, notes, result) = t
        values = [tt_id, "Selenium", fn, obj, inp, exp, tc_type,
                  ui, db, rb, notes, result]

        for col_idx, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col_idx, value=v)
            cell.border = BORDER
            cell.alignment = LEFT if col_idx in (3, 4, 5, 6, 11) else CENTER
            cell.font = NORMAL

        # TC Type column color (col 7)
        tc_cell = ws.cell(row=i, column=7)
        if tc_type == "Positive":
            tc_cell.fill = POSITIVE_FILL
            positive_count += 1
        elif tc_type == "Negative":
            tc_cell.fill = NEGATIVE_FILL
            negative_count += 1
        else:
            tc_cell.fill = MIXED_FILL
            mixed_count += 1
        tc_cell.font = BOLD

        # Result column color (col 12)
        rs_cell = ws.cell(row=i, column=12)
        if result == "PASS":
            rs_cell.fill = PASS_FILL
            pass_count += 1
        else:
            rs_cell.fill = FAIL_FILL
            fail_count += 1
        rs_cell.font = Font(bold=True, size=10,
                            color="006100" if result == "PASS" else "9C0006")

        ws.row_dimensions[i].height = 50

    # Footer summary
    footer_row = len(TESTS) + 4
    ws.merge_cells(start_row=footer_row, start_column=1,
                   end_row=footer_row, end_column=n_cols)
    total = len(TESTS)
    pct = round(pass_count / total * 100, 1)
    footer = ws.cell(row=footer_row, column=1,
                     value=f"Total: {total} TC  |  Positive: {positive_count}  |  "
                           f"Negative: {negative_count}  |  Mixed: {mixed_count}  |  "
                           f"Pass: {pass_count}/{total} ({pct}%)  |  "
                           f"Fail: {fail_count}")
    style_cell(footer, fill=SUBTITLE_FILL, font=Font(bold=True, size=11),
               alignment=CENTER)
    ws.row_dimensions[footer_row].height = 26

    # Run command note
    note_row = footer_row + 1
    ws.merge_cells(start_row=note_row, start_column=1,
                   end_row=note_row, end_column=n_cols)
    note = ws.cell(row=note_row, column=1,
                   value="Run command: pytest tests/ -v "
                         "--html=reports/selenium_report.html --self-contained-html")
    style_cell(note, font=Font(italic=True, size=9), alignment=CENTER)
    ws.row_dimensions[note_row].height = 18

    # Column widths
    widths = {
        1: 10, 2: 10, 3: 32, 4: 35, 5: 38, 6: 38,
        7: 11, 8: 13, 9: 13, 10: 16, 11: 38, 12: 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"Report generated: {OUT}")
    print(f"Total: {total} TC | Pass: {pass_count} | Fail: {fail_count}")


if __name__ == "__main__":
    main()
