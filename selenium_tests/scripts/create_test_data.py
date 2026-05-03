"""Tao file Excel chua du lieu test cho Selenium data-driven tests."""
import os
from openpyxl import Workbook

OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data",
    "auth_test_data.xlsx",
)

wb = Workbook()

# Sheet 1: StudentLogin
ws = wb.active
ws.title = "StudentLogin"
ws.append(["tc_id", "student_code", "password", "expected", "tc_type"])
ws.append(["TC001", "SV_TEST_01", "NewTest@12345", "success", "Positive"])
ws.append(["TC002", "SV_TEST_01", "wrong_pass", "khong dung", "Negative"])
ws.append(["TC003", "B99999_NOT_EXIST", "any_pass", "khong dung", "Negative"])
ws.append(["TC004", "", "pass", "validation", "Negative"])
ws.append(["TC005", "SV_TEST_01", "", "validation", "Negative"])

# Sheet 2: AdminLogin
ws2 = wb.create_sheet("AdminLogin")
ws2.append(["tc_id", "email", "password", "expected", "tc_type"])
ws2.append(["TC006", "admin@ptit.edu.vn", "admin123", "success", "Positive"])
ws2.append(["TC007", "notexist@ptit.edu.vn", "any", "Invalid credentials", "Negative"])
ws2.append(["TC008", "admin@ptit.edu.vn", "wrong", "Invalid credentials", "Negative"])

# Sheet 3: ChangePassword
ws3 = wb.create_sheet("ChangePassword")
ws3.append(["tc_id", "current", "new", "confirm", "expected", "tc_type"])
ws3.append(["TC013", "correct_pass", "NewPass123", "NewPass123", "success", "Positive"])
ws3.append(["TC014", "WRONG_PASS", "NewPass123", "NewPass123", "khong dung", "Negative"])
ws3.append(["TC016", "correct_pass", "abc", "abc", "it nhat 6", "Negative"])
ws3.append(["TC017", "correct_pass", "NewPass123", "Different99", "khong khop", "Negative"])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"Da tao: {OUT}")
