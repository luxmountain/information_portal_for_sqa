import os
import openpyxl
import pytest

from pages.login_page import LoginPage
from config.config import BASE_URL
from tests.accounts import STUDENT_CODE, STUDENT_PASS


def load_login_data(filepath=None, sheet="StudentLogin"):
    if filepath is None:
        filepath = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "data",
            "auth_test_data.xlsx",
        )
    if not os.path.exists(filepath):
        return []

    wb = openpyxl.load_workbook(filepath)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            data.append(
                {
                    "tc_id": row[0],
                    "code": row[1] or "",
                    "password": row[2] or "",
                    "expected": row[3] or "",
                    "tc_type": row[4] if len(row) > 4 else "",
                }
            )
    return data


TEST_DATA = load_login_data()


class TestStudentLogin:

    # TC001 - Dang nhap dung thong tin -> redirect khoi /login
    def test_TC001_valid_login(self, driver):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_student_code(STUDENT_CODE)
        page.enter_password(STUDENT_PASS)
        page.click_login()
        assert page.is_redirected_after_login(), (
            f"TC001 FAIL: Khong redirect sau dang nhap. URL hien tai: {driver.current_url}"
        )

    # TC002 - Sai mat khau
    def test_TC002_wrong_password(self, driver):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_student_code(STUDENT_CODE)
        page.enter_password("wrong_pass_xxx")
        page.click_login()
        error = page.get_error_message()
        assert error != "" and "/login" in driver.current_url, (
            f"TC002 FAIL: Khong hien loi. Error='{error}', URL={driver.current_url}"
        )

    # TC003 - Ma sinh vien khong ton tai
    def test_TC003_nonexistent_code(self, driver):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_student_code("B99999_NOT_EXIST")
        page.enter_password("any_pass")
        page.click_login()
        error = page.get_error_message()
        assert error != "" and "/login" in driver.current_url, (
            f"TC003 FAIL: Error='{error}', URL={driver.current_url}"
        )

    # TC004 - Bo trong ma sinh vien -> HTML5 required chan submit
    def test_TC004_missing_student_code(self, driver):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_password("any_pass")
        page.click_login()
        assert "/student/login" in driver.current_url, (
            "TC004 FAIL: Form da submit du thieu ma sinh vien"
        )

    # TC005 - Bo trong mat khau -> HTML5 required chan submit
    def test_TC005_missing_password(self, driver):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_student_code(STUDENT_CODE)
        page.click_login()
        assert "/student/login" in driver.current_url, (
            "TC005 FAIL: Form da submit du thieu mat khau"
        )

    # Data-driven test tu Excel
    @pytest.mark.skipif(not TEST_DATA, reason="Khong co du lieu test trong Excel")
    @pytest.mark.parametrize("td", TEST_DATA)
    def test_login_data_driven(self, driver, td):
        page = LoginPage(driver, BASE_URL)
        page.open_student_login()
        page.enter_student_code(td["code"])
        page.enter_password(td["password"])
        page.click_login()

        if td["expected"] == "success":
            assert page.is_redirected_after_login(), (
                f"{td['tc_id']} FAIL: Expected redirect sau dang nhap"
            )
        elif td["expected"] == "validation":
            # HTML5 required chan submit -> URL khong doi
            assert "/student/login" in driver.current_url, (
                f"{td['tc_id']} FAIL: Form submit du thieu input"
            )
        else:
            # Mong doi co loi tren UI
            error = page.get_error_message()
            assert error != "" and "/login" in driver.current_url, (
                f"{td['tc_id']} FAIL: Expected error containing '{td['expected']}', got '{error}'"
            )
