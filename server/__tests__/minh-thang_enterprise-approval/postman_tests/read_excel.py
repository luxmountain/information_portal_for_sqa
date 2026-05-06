import openpyxl

# Read TC091
wb = openpyxl.load_workbook('TC091_students_valid.xlsx')
ws = wb.active
print("=== TC091_students_valid.xlsx ===")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    if row[0] or row[1]:  # if has data
        print(f"Row {i}: {row}")
    else:
        break

# Read TC092
wb = openpyxl.load_workbook('TC092_students_partial.xlsx')
ws = wb.active
print("\n=== TC092_students_partial.xlsx ===")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    if row[0] or row[1]:
        print(f"Row {i}: {row}")
    else:
        break
